import asyncio
import csv
import io
from collections import Counter
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient
from langchain_core.runnables import RunnableLambda
from pydantic import ValidationError

from app.documents import DocumentExtractionError, extract_requirement_text
from app.main import app
from app.schemas import AcceptanceCriterion, AgileBacklog, NonFunctionalRequirement, UserStory
from app.service import (
    BacklogValidationError,
    GeminiRateLimitError,
    GeminiUnavailableError,
    generate_story,
)

client = TestClient(app)


def pdf_with_text(text: str) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects.append(
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    )
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(pdf)


def example_backlog() -> AgileBacklog:
    return AgileBacklog.model_validate(
        {
            "requirement_summary": {
                "business_objective": "Let users export a report.",
                "actors": ["user"],
                "key_capabilities": ["Export reports"],
                "constraints_and_business_rules": [],
            },
            "epics": [
                {
                    "epic_id": "E1",
                    "title": "Report export",
                    "description": "Export reports.",
                    "business_value": "Share reports.",
                    "related_story_ids": ["S1"],
                }
            ],
            "user_stories": [
                {
                    "story_id": "S1",
                    "epic_id": "E1",
                    "title": '=HYPERLINK("https://example.test")',
                    "persona": "user",
                    "user_story": "As a user, I want to export a report, so that I can share it.",
                    "acceptance_criteria": [
                        {
                            "given": "a report is available",
                            "when": "the user exports it",
                            "then": "a CSV file is downloaded",
                        }
                    ],
                    "suggested_priority": "Not suggested",
                    "priority_rationale": "The requirement does not specify priority.",
                    "dependencies": [],
                    "refinement_status": "Ready for Refinement Review",
                    "requirement_references": ["R-001"],
                    "assumptions": [],
                    "open_questions": ["Which report formats are required?"],
                }
            ],
            "non_functional_requirements": [
                {
                    "item_id": "NFR-001",
                    "requirement": "The export should be fast.",
                    "related_item_ids": ["S1"],
                    "testability_or_clarification": "No measurable threshold is specified.",
                }
            ],
            "assumptions": ["Unconfirmed: user means an authorized report user."],
            "open_questions": [
                {"question": "Which report formats are required?", "affected_items": ["S1"]}
            ],
            "out_of_scope": [],
            "requirement_coverage": [
                {
                    "requirement_id": "R-001",
                    "requirement_type": "Functional",
                    "explicit_requirement": "export a report",
                    "mapped_items": ["S1"],
                    "status": "Covered",
                    "reason": "S1 covers report export.",
                },
                {
                    "requirement_id": "R-002",
                    "requirement_type": "Non-functional",
                    "explicit_requirement": "The export should be fast.",
                    "mapped_items": ["NFR-001"],
                    "status": "Partially covered",
                    "reason": "A measurable performance threshold is not specified.",
                },
            ],
            "refinement_recommendations": [
                {"recommendation": "Confirm the supported report formats.", "related_items": ["S1"]}
            ],
        }
    )


def test_health() -> None:
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_swagger_openapi_documents_text_and_file_generation() -> None:
    schema = client.get("/openapi.json").json()
    paths = schema["paths"]

    text_operation = paths["/api/v1/stories/generate"]["post"]
    file_operation = paths["/api/v1/stories/generate-file"]["post"]
    export_operation = paths["/api/v1/stories/export"]["post"]

    assert text_operation["tags"] == ["Generate backlog"]
    assert text_operation["requestBody"]["content"]["application/json"]["schema"]["$ref"]
    request_schema = schema["components"]["schemas"]["GenerateStoryRequest"]
    assert request_schema["examples"][0]["requirement"]
    assert "\\n" in request_schema["examples"][0]["requirement"]
    assert file_operation["tags"] == ["Generate backlog"]
    assert "multipart/form-data" in file_operation["requestBody"]["content"]
    file_schema = schema["components"]["schemas"][
        "Body_create_story_from_file_api_v1_stories_generate_file_post"
    ]
    assert "Supported extensions" in file_schema["properties"]["file"]["description"]
    assert export_operation["tags"] == ["Export"]
    assert "complete JSON response" in export_operation["description"]


def test_generation_preflight_allows_local_vite_origins() -> None:
    for origin in ("http://localhost:5173", "http://127.0.0.1:5173"):
        response = client.options(
            "/api/v1/stories/generate",
            headers={
                "Origin": origin,
                "Access-Control-Request-Method": "POST",
                "Access-Control-Request-Headers": "content-type",
            },
        )
        assert response.status_code == 200
        assert response.headers["access-control-allow-origin"] == origin


def test_generation_returns_validated_backlog(monkeypatch) -> None:
    async def fake_generate(requirement: str) -> AgileBacklog:
        assert requirement == "Users can export reports as CSV files."
        return example_backlog()

    monkeypatch.setattr("app.main.generate_story", fake_generate)
    response = client.post(
        "/api/v1/stories/generate",
        json={"requirement": "Users can export reports as CSV files."},
    )
    assert response.status_code == 200
    assert response.json()["backlog"]["user_stories"][0]["story_id"] == "S1"
    assert response.json()["backlog"]["requirement_summary"]["business_objective"]
    assert response.json()["source_requirement"] == "Users can export reports as CSV files."


def test_generation_retries_invalid_references_once(monkeypatch) -> None:
    invalid_data = example_backlog().model_dump()
    invalid_data["user_stories"][0]["requirement_references"] = ["R-009"]
    with pytest.raises(ValidationError) as validation_error:
        AgileBacklog.model_validate(invalid_data)

    responses = [
        {"raw": None, "parsed": None, "parsing_error": validation_error.value},
        {"raw": None, "parsed": example_backlog(), "parsing_error": None},
    ]

    class FakeModel:
        def with_structured_output(self, _schema, **kwargs):
            assert kwargs == {"method": "json_schema", "include_raw": True}
            return RunnableLambda(lambda _input: responses.pop(0))

    monkeypatch.setattr(
        "app.service.get_settings",
        lambda: SimpleNamespace(google_api_key="test-key", gemini_model="test-model"),
    )
    monkeypatch.setattr("app.service.ChatGoogleGenerativeAI", lambda **_kwargs: FakeModel())

    result = asyncio.run(generate_story("A valid product requirement for backlog generation."))

    assert result == example_backlog()
    assert responses == []


def test_generation_raises_clear_error_after_invalid_retry(monkeypatch) -> None:
    invalid_data = example_backlog().model_dump()
    invalid_data["user_stories"][0]["requirement_references"] = ["R-009"]
    with pytest.raises(ValidationError) as validation_error:
        AgileBacklog.model_validate(invalid_data)

    responses = [
        {"raw": None, "parsed": None, "parsing_error": validation_error.value},
        {"raw": None, "parsed": None, "parsing_error": validation_error.value},
    ]

    class FakeModel:
        def with_structured_output(self, _schema, **_kwargs):
            return RunnableLambda(lambda _input: responses.pop(0))

    monkeypatch.setattr(
        "app.service.get_settings",
        lambda: SimpleNamespace(google_api_key="test-key", gemini_model="test-model"),
    )
    monkeypatch.setattr("app.service.ChatGoogleGenerativeAI", lambda **_kwargs: FakeModel())

    with pytest.raises(BacklogValidationError, match="after one automatic correction attempt"):
        asyncio.run(generate_story("A valid product requirement for backlog generation."))
    assert responses == []


def test_generation_returns_actionable_error_if_corrective_retry_fails(monkeypatch) -> None:
    async def fail_validation(_requirement: str) -> AgileBacklog:
        raise BacklogValidationError(
            "Gemini returned an inconsistent backlog after one automatic correction attempt. "
            "Please retry the request."
        )

    monkeypatch.setattr("app.main.generate_story", fail_validation)
    response = client.post(
        "/api/v1/stories/generate",
        json={"requirement": "A valid product requirement for backlog generation."},
    )

    assert response.status_code == 502
    assert "after one automatic correction attempt" in response.json()["detail"]


def test_generation_returns_rate_limit_status_and_retry_after(monkeypatch) -> None:
    async def fail_rate_limit(_requirement: str) -> AgileBacklog:
        raise GeminiRateLimitError("Gemini rate limit is temporarily unavailable.", retry_after=7)

    monkeypatch.setattr("app.main.generate_story", fail_rate_limit)
    response = client.post(
        "/api/v1/stories/generate",
        json={"requirement": "A valid product requirement for backlog generation."},
    )

    assert response.status_code == 429
    assert response.headers["retry-after"] == "7"
    assert "rate limit" in response.json()["detail"]


def test_generation_returns_service_unavailable_after_transient_failures(monkeypatch) -> None:
    async def fail_transient(_requirement: str) -> AgileBacklog:
        raise GeminiUnavailableError("Gemini is temporarily unavailable.")

    monkeypatch.setattr("app.main.generate_story", fail_transient)
    response = client.post(
        "/api/v1/stories/generate",
        json={"requirement": "A valid product requirement for backlog generation."},
    )

    assert response.status_code == 503
    assert "temporarily unavailable" in response.json()["detail"]


def test_multiline_swagger_requirement_accepts_escaped_json_newlines(monkeypatch) -> None:
    requirement = (
        "A food delivery company wants customers to browse restaurants, order food, make "
        "payments, and track deliveries.\n"
        "Restaurants manage orders, and delivery partners update delivery status.\n"
        "Business goal: Improve customer satisfaction and delivery efficiency."
    )

    async def fake_generate(value: str) -> AgileBacklog:
        assert value == requirement
        return example_backlog()

    monkeypatch.setattr("app.main.generate_story", fake_generate)
    response = client.post("/api/v1/stories/generate", json={"requirement": requirement})

    assert response.status_code == 200
    assert response.json()["source_requirement"] == requirement


def test_generation_explains_unescaped_newline_in_json_string() -> None:
    raw_json = '{"requirement": "Line one\nLine two, with enough text."}'
    response = client.post(
        "/api/v1/stories/generate",
        content=raw_json,
        headers={"Content-Type": "application/json"},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == (
        "Invalid JSON request body. Escape line breaks inside the requirement "
        "string as \\n, or enter the requirement on one line."
    )


def test_generation_rejects_short_requirement() -> None:
    response = client.post("/api/v1/stories/generate", json={"requirement": "short"})
    assert response.status_code == 422


def test_export_returns_full_backlog_csv_and_neutralizes_formulas() -> None:
    backlog = example_backlog()
    backlog.requirement_summary.business_objective = (
        'Help users export reports, including "quoted" report titles — with traceability.'
    )
    backlog.requirement_summary.actors = ["user", "analyst, QA"]
    backlog.requirement_summary.key_capabilities = ["Export reports", "Review coverage"]
    backlog.requirement_summary.constraints_and_business_rules = [
        "Preserve Unicode: café, 東京, and 🚀"
    ]
    backlog.user_stories[0].title = '\x00\t=HYPERLINK("https://example.test")'
    backlog.user_stories[0].acceptance_criteria.append(
        AcceptanceCriterion(
            given="the requirement contains commas, quotes, and\nline breaks",
            when="the backlog is exported",
            then="all criterion fields remain intact",
        )
    )
    backlog.user_stories[0].assumptions = [
        "Unconfirmed: support UTF-8 titles.",
        "Unconfirmed: preserve commas, quotes, and line breaks.",
    ]
    second_story_data = backlog.user_stories[0].model_dump()
    second_story_data.update(
        {
            "story_id": "S2",
            "title": "Review export coverage",
            "epic_id": "E1",
            "dependencies": ["S1"],
            "requirement_references": ["R-001", "R-002"],
            "acceptance_criteria": [
                {
                    "given": "a backlog has been generated",
                    "when": "the analyst reviews its coverage",
                    "then": "all requirement mappings are displayed",
                }
            ],
            "assumptions": [],
            "open_questions": [],
        }
    )
    backlog.user_stories.append(UserStory.model_validate(second_story_data))
    backlog.assumptions.append(" \t+SUM(1,2)")
    backlog.out_of_scope = ["Exporting to PDF is excluded."]
    response = client.post(
        "/api/v1/stories/export",
        json={
            "source_requirement": (
                '=IMPORTDATA("https://example.test"),\nwith Unicode: café, 東京, and 🚀'
            ),
            "backlog": backlog.model_dump(),
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert 'attachment; filename="agile-backlog.csv"' == response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    expected_headers = [
        "Record Type",
        "ID",
        "Title",
        "Description",
        "Related Item IDs",
        "Source Requirement",
        "Business Objective",
        "Actors",
        "Key Capabilities",
        "Business Rules and Constraints",
        "Business Value",
        "Related Story IDs",
        "Epic ID",
        "Persona",
        "User Story",
        "Acceptance Criteria",
        "Priority",
        "Priority Rationale",
        "Refinement Status",
        "Dependencies",
        "Requirement References",
        "Assumptions",
        "Open Questions",
        "Requirement Text",
        "Related Items",
        "Testability or Clarification",
        "Affected Items",
        "Out-of-Scope Items",
        "Requirement Type",
        "Explicit Requirement",
        "Mapped Backlog Items",
        "Coverage Status",
        "Coverage Reason",
        "Recommendation",
    ]
    assert list(rows[0]) == expected_headers
    assert all(len(row) == len(expected_headers) for row in rows)
    assert Counter(row["Record Type"] for row in rows) == {
        "Source Requirement": 1,
        "Requirement Summary": 1,
        "Epic": 1,
        "User Story": 2,
        "Non-Functional Requirement": 1,
        "Assumption": 2,
        "Open Question": 1,
        "Out of Scope": 1,
        "Requirement Coverage": 2,
        "Refinement Recommendation": 1,
    }

    def item(section: str, item_id: str) -> dict[str, str]:
        return next(row for row in rows if row["Record Type"] == section and row["ID"] == item_id)

    source_row = item("Source Requirement", "SOURCE")
    assert source_row["Source Requirement"].startswith("'=")
    assert "café, 東京, and 🚀" in source_row["Source Requirement"]

    summary = item("Requirement Summary", "SUMMARY")
    assert summary["Business Objective"] == (
        'Help users export reports, including "quoted" report titles — with traceability.'
    )
    assert summary["Actors"] == "user; analyst, QA"
    assert summary["Key Capabilities"] == "Export reports; Review coverage"
    assert summary["Business Rules and Constraints"] == "Preserve Unicode: café, 東京, and 🚀"

    epic = item("Epic", "E1")
    assert epic["Title"] == "Report export"
    assert epic["Description"] == "Export reports."
    assert epic["Business Value"] == "Share reports."
    assert epic["Related Story IDs"] == "S1; S2"

    story = item("User Story", "S1")
    assert story["Title"].startswith("'\x00\t=")
    assert story["Epic ID"] == "E1"
    assert story["Related Item IDs"] == "E1"
    assert story["Persona"] == "user"
    assert story["User Story"] == ("As a user, I want to export a report, so that I can share it.")
    assert story["Priority"] == "Not suggested"
    assert story["Priority Rationale"] == "The requirement does not specify priority."
    assert story["Refinement Status"] == "Ready for Refinement Review"
    assert story["Dependencies"] == ""
    assert story["Requirement References"] == "R-001"
    assert story["Assumptions"] == (
        "Unconfirmed: support UTF-8 titles.\nUnconfirmed: preserve commas, quotes, and line breaks."
    )
    assert story["Open Questions"] == "Which report formats are required?"
    criteria = story["Acceptance Criteria"]
    assert criteria == (
        "Given: a report is available\n"
        "When: the user exports it\n"
        "Then: a CSV file is downloaded\n\n"
        "Given: the requirement contains commas, quotes, and\n"
        "line breaks\n"
        "When: the backlog is exported\n"
        "Then: all criterion fields remain intact"
    )
    assert len([row for row in rows if row["Record Type"] == "Acceptance Criterion"]) == 0

    second_story = item("User Story", "S2")
    assert second_story["Dependencies"] == "S1"
    assert second_story["Requirement References"] == "R-001; R-002"
    assert second_story["Acceptance Criteria"].startswith("Given: a backlog has been generated")

    nfr = item("Non-Functional Requirement", "NFR-001")
    assert nfr["Requirement Text"] == "The export should be fast."
    assert nfr["Related Items"] == nfr["Related Item IDs"] == "S1"
    assert nfr["Testability or Clarification"] == "No measurable threshold is specified."
    assert item("Assumption", "ASSUMPTION-001")["Assumptions"] == (
        "Unconfirmed: user means an authorized report user."
    )
    assert item("Assumption", "ASSUMPTION-002")["Assumptions"] == "' \t+SUM(1,2)"
    question = item("Open Question", "QUESTION-001")
    assert question["Open Questions"] == "Which report formats are required?"
    assert question["Affected Items"] == question["Related Item IDs"] == "S1"
    assert item("Out of Scope", "OUT-001")["Out-of-Scope Items"] == (
        "Exporting to PDF is excluded."
    )
    coverage = item("Requirement Coverage", "R-001")
    assert coverage["Requirement Type"] == "Functional"
    assert coverage["Explicit Requirement"] == "export a report"
    assert coverage["Mapped Backlog Items"] == coverage["Related Item IDs"] == "S1"
    assert coverage["Coverage Status"] == "Covered"
    assert coverage["Coverage Reason"] == "S1 covers report export."
    recommendation = item("Refinement Recommendation", "RECOMMENDATION-001")
    assert recommendation["Recommendation"] == "Confirm the supported report formats."
    assert recommendation["Related Item IDs"] == "S1"


def test_export_handles_empty_optional_values_without_inventing_rows() -> None:
    backlog = example_backlog()
    backlog.requirement_summary.actors = []
    backlog.requirement_summary.key_capabilities = []
    backlog.requirement_summary.constraints_and_business_rules = []
    backlog.user_stories[0].epic_id = None
    backlog.user_stories[0].dependencies = []
    backlog.user_stories[0].assumptions = []
    backlog.user_stories[0].open_questions = []
    backlog.non_functional_requirements = []
    backlog.requirement_coverage[1].mapped_items = []
    backlog.assumptions = []
    backlog.open_questions = []
    backlog.out_of_scope = []
    backlog.refinement_recommendations = []

    response = client.post(
        "/api/v1/stories/export",
        json={
            "source_requirement": "A sufficiently long requirement with optional data omitted.",
            "backlog": backlog.model_dump(),
        },
    )

    assert response.status_code == 200
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0]["Actors"] == ""
    assert rows[0]["Key Capabilities"] == ""
    assert rows[0]["Business Rules and Constraints"] == ""
    story = next(row for row in rows if row["Record Type"] == "User Story")
    assert story["Epic ID"] == ""
    assert story["Related Item IDs"] == ""
    assert story["Dependencies"] == ""
    assert story["Assumptions"] == ""
    assert story["Open Questions"] == ""
    coverage = next(
        row for row in rows if row["Record Type"] == "Requirement Coverage" and row["ID"] == "R-002"
    )
    assert coverage["Mapped Backlog Items"] == ""
    assert Counter(row["Record Type"] for row in rows) == {
        "Source Requirement": 1,
        "Requirement Summary": 1,
        "Epic": 1,
        "User Story": 1,
        "Requirement Coverage": 2,
    }


def test_normalized_csv_remains_available_as_an_optional_export() -> None:
    response = client.post(
        "/api/v1/stories/export?format=normalized",
        json={
            "source_requirement": "A sufficiently long requirement for export.",
            "backlog": example_backlog().model_dump(),
        },
    )

    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        'attachment; filename="agile-backlog-normalized.csv"'
    )
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert list(rows[0]) == ["Section", "Item ID", "Related Item IDs", "Field", "Value"]
    story_title = next(
        row
        for row in rows
        if row["Section"] == "User Story" and row["Item ID"] == "S1" and row["Field"] == "Title"
    )
    assert story_title["Value"].startswith("'=")
    criteria = next(
        row
        for row in rows
        if row["Section"] == "User Story"
        and row["Item ID"] == "S1"
        and row["Field"] == "Acceptance Criteria"
    )
    assert "Given: a report is available" in criteria["Value"]


def test_backlog_rejects_dangling_requirement_reference() -> None:
    data = example_backlog().model_dump()
    data["user_stories"][0]["requirement_references"] = ["R-999"]
    with pytest.raises(ValueError, match="unknown IDs: R-999"):
        AgileBacklog.model_validate(data)


def test_backlog_repairs_epic_story_membership_from_story_epic_id() -> None:
    data = example_backlog().model_dump()
    data["epics"][0]["related_story_ids"] = []

    backlog = AgileBacklog.model_validate(data)

    assert backlog.user_stories[0].epic_id == "E1"
    assert backlog.epics[0].related_story_ids == ["S1"]


def test_backlog_rejects_story_with_unknown_epic() -> None:
    data = example_backlog().model_dump()
    data["user_stories"][0]["epic_id"] = "E9"
    data["epics"][0]["related_story_ids"] = []

    with pytest.raises(ValueError, match="S1 references unknown epic E9"):
        AgileBacklog.model_validate(data)


def test_non_functional_item_does_not_duplicate_requirement_id() -> None:
    item = {
        "item_id": "NFR-001",
        "requirement": "Export response has a stated performance target.",
        "related_item_ids": ["S1"],
        "testability_or_clarification": "Define a measurable threshold.",
    }

    parsed = NonFunctionalRequirement.model_validate({**item, "requirement_id": "R-NFR-01"})

    assert parsed.item_id == "NFR-001"
    assert "requirement_id" not in parsed.model_dump()


def test_backlog_ignores_legacy_nfr_requirement_id() -> None:
    data = example_backlog().model_dump()
    data["non_functional_requirements"][0]["requirement_id"] = "R-NFR-01"

    backlog = AgileBacklog.model_validate(data)

    assert backlog.non_functional_requirements[0].item_id == "NFR-001"
    assert "requirement_id" not in backlog.non_functional_requirements[0].model_dump()


def test_generate_from_txt_upload(monkeypatch) -> None:
    requirement = "Users can export reports as CSV files."

    async def fake_generate(value: str) -> AgileBacklog:
        assert value == requirement
        return example_backlog()

    monkeypatch.setattr("app.main.generate_story", fake_generate)
    response = client.post(
        "/api/v1/stories/generate-file",
        files={"file": ("requirements.txt", requirement.encode(), "text/plain")},
    )
    assert response.status_code == 200
    assert response.json()["source_requirement"] == requirement


def test_generate_file_rejects_upload_over_configured_limit(monkeypatch) -> None:
    monkeypatch.setattr("app.main.settings.max_upload_bytes", 16)
    response = client.post(
        "/api/v1/stories/generate-file",
        files={
            "file": ("requirements.txt", b"this file exceeds the configured limit", "text/plain")
        },
    )
    assert response.status_code == 413
    assert "upload limit" in response.json()["detail"]


def test_generate_from_docx_upload(monkeypatch) -> None:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Users can export reports as CSV files.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Reports include their title."
    document.save(buffer)
    expected = "Users can export reports as CSV files.\nReports include their title."

    async def fake_generate(value: str) -> AgileBacklog:
        assert value == expected
        return example_backlog()

    monkeypatch.setattr("app.main.generate_story", fake_generate)
    response = client.post(
        "/api/v1/stories/generate-file",
        files={
            "file": (
                "requirements.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    assert response.json()["source_requirement"] == expected


def test_generate_from_pdf_upload(monkeypatch) -> None:
    requirement = "Users can export reports as CSV files."

    async def fake_generate(value: str) -> AgileBacklog:
        assert value == requirement
        return example_backlog()

    monkeypatch.setattr("app.main.generate_story", fake_generate)
    response = client.post(
        "/api/v1/stories/generate-file",
        files={"file": ("requirements.pdf", pdf_with_text(requirement), "application/pdf")},
    )
    assert response.status_code == 200
    assert response.json()["source_requirement"] == requirement


@pytest.mark.parametrize(
    ("filename", "content", "error"),
    [
        ("requirements.exe", b"plain text requirement", "Unsupported file type"),
        ("requirements.txt", b"\xff\xfe", "UTF-8"),
        ("requirements.pdf", b"not a pdf document", "not a valid PDF"),
        ("requirements.docx", b"not a word document", "Could not read the DOCX"),
    ],
)
def test_reject_invalid_uploads(filename: str, content: bytes, error: str) -> None:
    response = client.post(
        "/api/v1/stories/generate-file",
        files={"file": (filename, content, "application/octet-stream")},
    )
    assert response.status_code == 400
    assert error in response.json()["detail"]


def test_extract_requirement_rejects_empty_text_and_oversize_content() -> None:
    with pytest.raises(DocumentExtractionError, match="at least 10 characters"):
        extract_requirement_text("requirements.txt", b"short", 100)
    with pytest.raises(ValueError, match="exceeds the 10-character limit"):
        extract_requirement_text("requirements.txt", b"long requirement", 10)


def test_extract_pdf_rejects_image_only_document(monkeypatch) -> None:
    class ImageOnlyPdf:
        is_encrypted = False
        pages: list[object] = []

    monkeypatch.setattr("app.documents.PdfReader", lambda *_args, **_kwargs: ImageOnlyPdf())
    with pytest.raises(DocumentExtractionError, match="No readable text"):
        extract_requirement_text("requirements.pdf", b"%PDF-1.7", 10_000)
