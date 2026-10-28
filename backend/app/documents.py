import io
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class DocumentExtractionError(ValueError):
    pass


class RequirementTooLongError(ValueError):
    pass


def extract_requirement_text(filename: str, content: bytes, max_chars: int) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in {".txt", ".pdf", ".docx"}:
        raise DocumentExtractionError("Unsupported file type. Upload a .txt, .pdf, or .docx file.")
    if not content:
        raise DocumentExtractionError("The uploaded file is empty.")

    if extension == ".txt":
        if b"\x00" in content:
            raise DocumentExtractionError("The .txt file contains binary data, not plain text.")
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentExtractionError("The .txt file must use UTF-8 text encoding.") from exc
    elif extension == ".pdf":
        text = _extract_pdf(content, max_chars)
    else:
        text = _extract_docx(content, max_chars)

    text = text.strip()
    if not text:
        raise DocumentExtractionError(
            "No readable text was found in the document. Scanned PDFs need OCR before upload."
        )
    if len(text) > max_chars:
        raise RequirementTooLongError(
            f"Extracted requirement exceeds the {max_chars:,}-character limit."
        )
    if len(text) < 10:
        raise DocumentExtractionError("The document must contain at least 10 characters of text.")
    return text


def _extract_pdf(content: bytes, max_chars: int) -> str:
    if not content.startswith(b"%PDF-"):
        raise DocumentExtractionError("The uploaded file is not a valid PDF.")
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise DocumentExtractionError("Password-protected PDFs are not supported.")
        pages: list[str] = []
        total_chars = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            total_chars += len(page_text)
            if total_chars > max_chars:
                raise RequirementTooLongError(
                    f"Extracted requirement exceeds the {max_chars:,}-character limit."
                )
            if page_text:
                pages.append(page_text)
        return "\n\n".join(pages)
    except (DocumentExtractionError, RequirementTooLongError):
        raise
    except Exception as exc:
        raise DocumentExtractionError(
            "Could not read the PDF. Check that the file is valid."
        ) from exc


def _extract_docx(content: bytes, max_chars: int) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1_000 or sum(entry.file_size for entry in entries) > 30_000_000:
                raise DocumentExtractionError("The DOCX file contains too much expanded data.")
            filenames = archive.namelist()
            if "[Content_Types].xml" not in filenames or "word/document.xml" not in filenames:
                raise DocumentExtractionError("The uploaded file is not a valid DOCX document.")
        document = Document(io.BytesIO(content))
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(
            "Could not read the DOCX file. Check that the document is valid."
        ) from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if len(text) > max_chars:
        raise RequirementTooLongError(
            f"Extracted requirement exceeds the {max_chars:,}-character limit."
        )
    return text
